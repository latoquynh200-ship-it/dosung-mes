import streamlit as st
import pandas as pd
import openpyxl
import os
import calendar
from datetime import datetime

st.set_page_config(
    page_title="DO SUNG MES - Delivery Management Hub",
    page_icon="🏭",
    layout="wide",
    initial_sidebar_state="collapsed"
)

st.markdown("""
<style>
    .main { background-color: #0e1117; }
    .metric-card {
        background: linear-gradient(135deg, #1e2530 0%, #151a21 100%);
        border: 1px solid #2d3748;
        padding: 16px 20px;
        border-radius: 10px;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.2);
    }
    .metric-title { color: #a0aec0; font-size: 13px; font-weight: 500; text-transform: uppercase; letter-spacing: 0.5px; }
    .metric-value { color: #f7fafc; font-size: 26px; font-weight: 700; margin-top: 4px; }
    .metric-sub { color: #718096; font-size: 12px; margin-top: 4px; }
    .stButton>button {
        background-color: #2563eb;
        color: white;
        font-weight: 600;
        border-radius: 6px;
        border: none;
        padding: 8px 20px;
    }
</style>
""", unsafe_allow_html=True)

# Xác thực truy cập
if "authenticated" not in st.session_state:
    st.session_state.authenticated = False

if not st.session_state.authenticated:
    c_space1, c_login, c_space2 = st.columns([1, 1.2, 1])
    with c_login:
        st.markdown("<br><br>", unsafe_allow_html=True)
        st.markdown("""
        <div class="metric-card" style="text-align: center; padding: 30px;">
            <h2 style="color: #f8fafc; margin-bottom: 5px;">🏭 DO SUNG MES PORTAL</h2>
            <p style="color: #94a3b8; font-size: 14px;">Hệ Thống Theo Dõi Giao Nhận & Kế Hoạch 2026</p>
        </div>
        """, unsafe_allow_html=True)
        
        with st.form("auth_form"):
            pwd = st.text_input("Mã xác thực đối tác / Access Key:", type="password")
            submit_auth = st.form_submit_button("Đăng Nhập Vào Hệ Thống", use_container_width=True)
            if submit_auth:
                if pwd == "dosung2026":
                    st.session_state.authenticated = True
                    st.rerun()
                else:
                    st.warning("⚠️ Mã truy cập không chính xác.")
    st.stop()

# Quét tìm file dữ liệu
current_dir = os.path.dirname(os.path.abspath(__file__))
excel_candidates = [
    os.path.join(current_dir, f) for f in os.listdir(current_dir)
    if (".xlsx" in f.lower() or ".xls" in f.lower()) and not f.startswith("~$")
]
target_file = excel_candidates[0] if excel_candidates else None

header_l, header_r = st.columns([3, 1])
with header_l:
    st.title("🏭 QUẢN LÝ LỊCH GIAO NHẬN & TIẾN ĐỘ DO SUNG")
    file_name = os.path.basename(target_file) if target_file else "Chưa tìm thấy"
    st.caption(f"Trạng thái: 🟢 **Trực tuyến** | Nguồn dữ liệu: `{file_name}`")
with header_r:
    if st.button("Đăng Xuất", use_container_width=True):
        st.session_state.authenticated = False
        st.rerun()

if not target_file:
    st.info("💡 Chưa tìm thấy file dữ liệu. Vui lòng tải file Excel lên:")
    uploaded = st.file_uploader("Kéo và thả file Excel (.xlsx):", type=["xlsx", "xls"])
    if uploaded:
        save_path = os.path.join(current_dir, "dosung_schedule.xlsx")
        with open(save_path, "wb") as f:
            f.write(uploaded.getbuffer())
        st.rerun()
    st.stop()

try:
    wb = openpyxl.load_workbook(target_file, data_only=True)
    available_sheets = [s for s in wb.sheetnames if s != "Tổng nhập"]
except Exception:
    st.warning("Không thể đọc tệp Excel. Vui lòng kiểm tra định dạng.")
    st.stop()

tab_daily, tab_summary, tab_export = st.tabs([
    "📅 Lịch Giao Nhận & Cập Nhật Ngày", 
    "📊 Đối Soát Tổng Nhập & Hải Quan", 
    "📥 Trích Xuất Dữ Liệu"
])

# ==================== TAB 1 ====================
with tab_daily:
    col1, col2 = st.columns([1, 1])
    with col1:
        selected_part = st.selectbox("📌 Mã linh kiện / Sheet:", available_sheets, index=0)
    
    sheet = wb[selected_part]

    month_dict = {}
    for r in range(1, min(sheet.max_row + 1, 120)):
        v_year = sheet.cell(r, 2).value
        v_month = sheet.cell(r, 5).value
        if v_year == 2026 and v_month is not None:
            try:
                month_dict[int(v_month)] = r
            except:
                pass

    with col2:
        if month_dict:
            month_opts = sorted(list(month_dict.keys()))
            selected_month_num = st.selectbox(
                "🗓️ Tháng vận hành (Năm 2026):", 
                month_opts, 
                format_func=lambda x: f"Tháng {x:02d}"
            )
            month_start_row = month_dict[selected_month_num]
        else:
            selected_month_num = 1
            month_start_row = 5

    row_po = month_start_row + 3
    row_plan = month_start_row + 4
    row_actual = month_start_row + 5
    row_ng = month_start_row + 6

    def safe_num(val):
        if val is None: return 0
        try: return int(val)
        except: return 0

    num_days = calendar.monthrange(2026, selected_month_num)[1]

    days_data = []
    for day in range(1, num_days + 1):
        col = 6 + day
        p_val = safe_num(sheet.cell(row=row_plan, column=col).value)
        a_val = safe_num(sheet.cell(row=row_actual, column=col).value)
        n_val = safe_num(sheet.cell(row=row_ng, column=col).value)
        po_val = safe_num(sheet.cell(row=row_po, column=col).value)

        days_data.append({
            "Ngày": f"Ngày {day:02d}",
            "Day_Num": day,
            "PO": po_val,
            "Kế Hoạch": p_val,
            "Thực Nhập": a_val,
            "Lỗi (NG)": n_val,
            "Chênh Lệch": a_val - p_val
        })

    df_month = pd.DataFrame(days_data)

    tot_plan = int(df_month["Kế Hoạch"].sum())
    tot_actual = int(df_month["Thực Nhập"].sum())
    tot_ng = int(df_month["Lỗi (NG)"].sum())
    completion_rate = (tot_actual / tot_plan * 100) if tot_plan > 0 else 0.0

    kp1, kp2, kp3, kp4 = st.columns(4)
    with kp1:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-title">Tổng Kế Hoạch Tháng</div>
            <div class="metric-value">{tot_plan:,.0f} <span style="font-size: 14px; font-weight: normal; color: #94a3b8;">PCS</span></div>
            <div class="metric-sub">Kế hoạch chỉ định</div>
        </div>
        """, unsafe_allow_html=True)
    with kp2:
        diff_color = "#34d399" if (tot_actual - tot_plan) >= 0 else "#fbbf24"
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-title">Thực Nhập Lũy Kế</div>
            <div class="metric-value">{tot_actual:,.0f} <span style="font-size: 14px; font-weight: normal; color: #94a3b8;">PCS</span></div>
            <div class="metric-sub" style="color: {diff_color};">Độ lệch: {tot_actual - tot_plan:+,.0f} PCS</div>
        </div>
        """, unsafe_allow_html=True)
    with kp3:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-title">Phát Sinh Hàng Lỗi (NG)</div>
            <div class="metric-value" style="color: #f87171;">{tot_ng:,.0f} <span style="font-size: 14px; font-weight: normal; color: #94a3b8;">PCS</span></div>
            <div class="metric-sub">Số lượng lỗi / sửa</div>
        </div>
        """, unsafe_allow_html=True)
    with kp4:
        st.markdown(f"""
        <div class="metric-card">
            <div class="metric-title">Tỷ Lệ Hoàn Thành</div>
            <div class="metric-value" style="color: #60a5fa;">{completion_rate:.1f}%</div>
            <div class="metric-sub">Tiến độ giao hàng</div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    f_col, t_col = st.columns([1, 2.2])

    with f_col:
        st.markdown("""
        <div style="background: #151a21; border: 1px solid #2d3748; padding: 18px; border-radius: 8px;">
            <h4 style="margin: 0 0 12px 0; color: #f1f5f9;">📝 Ghi Nhận Thực Tế</h4>
        """, unsafe_allow_html=True)
        
        with st.form("entry_form"):
            in_day = st.number_input("Chọn ngày trong tháng:", min_value=1, max_value=num_days, value=1)
            
            default_act = 0
            default_ng = 0
            if in_day in df_month["Day_Num"].values:
                default_act = int(df_month.loc[df_month["Day_Num"] == in_day, "Thực Nhập"].values[0])
                default_ng = int(df_month.loc[df_month["Day_Num"] == in_day, "Lỗi (NG)"].values[0])

            in_act = st.number_input("Số lượng Thực Nhập (PCS):", min_value=0, step=50, value=default_act)
            in_ng = st.number_input("Số lượng Lỗi NG (PCS):", min_value=0, step=1, value=default_ng)
            
            st.write("")
            btn_save = st.form_submit_button("Cập Nhật Dữ Liệu", use_container_width=True)

            if btn_save:
                try:
                    wb_w = openpyxl.load_workbook(target_file)
                    ws_w = wb_w[selected_part]
                    t_col_idx = 6 + int(in_day)
                    ws_w.cell(row=row_actual, column=t_col_idx, value=int(in_act))
                    ws_w.cell(row=row_ng, column=t_col_idx, value=int(in_ng))
                    wb_w.save(target_file)
                    st.toast(f"Đã lưu thành công Ngày {in_day:02d}!", icon="✅")
                    st.rerun()
                except Exception:
                    st.warning("Tệp đang bận, vui lòng thử lại.")
        
        st.markdown("</div>", unsafe_allow_html=True)

    with t_col:
        st.markdown(f"<h4 style='color: #f1f5f9; margin-bottom: 12px;'>📋 Chi Tiết Lịch Giao Nhận — Tháng {selected_month_num:02d} ({selected_part})</h4>", unsafe_allow_html=True)
        
        def style_variance(val):
            if val < 0:
                return 'color: #fbbf24; font-weight: 600;'
            elif val > 0:
                return 'color: #34d399; font-weight: 600;'
            return 'color: #94a3b8;'

        view_df = df_month[["Ngày", "PO", "Kế Hoạch", "Thực Nhập", "Lỗi (NG)", "Chênh Lệch"]]
        st.dataframe(
            view_df.style.map(style_variance, subset=['Chênh Lệch']),
            use_container_width=True,
            height=370
        )

# ==================== TAB 2 ====================
with tab_summary:
    st.markdown("<h4 style='color: #f1f5f9;'>📑 Báo Cáo Đối Soát Tổng Nhập & Hồ Sơ Hải Quan</h4>", unsafe_allow_html=True)
    if "Tổng nhập" in wb.sheetnames:
        ws_tn = wb["Tổng nhập"]
        summary_rows = []
        for r in range(4, 12):
            m_label = ws_tn.cell(row=r, column=2).value
            if m_label is not None:
                summary_rows.append({
                    "Kỳ Tháng": str(m_label),
                    "PAR-011: Đã Giao": ws_tn.cell(row=r, column=3).value or 0,
                    "PAR-011: Khai HQ": ws_tn.cell(row=r, column=4).value or 0,
                    "PAR-011: Tồn Chưa Khai": ws_tn.cell(row=r, column=5).value or 0,
                    "Woodone: Đã Giao": ws_tn.cell(row=r, column=6).value or 0,
                    "Woodone: Khai HQ": ws_tn.cell(row=r, column=7).value or 0,
                    "mda0016: Đã Giao": ws_tn.cell(row=r, column=9).value or 0,
                })
        st.dataframe(pd.DataFrame(summary_rows), use_container_width=True)

        st.markdown("<br>", unsafe_allow_html=True)
        po_c1, po_c2 = st.columns([1.2, 1])
        with po_c1:
            st.markdown("<h5 style='color: #cbd5e1;'>📦 Danh Mục Đơn Đặt Hàng (PO Tracking)</h5>", unsafe_allow_html=True)
            po_items = []
            for r in range(4, 11):
                p_code = ws_tn.cell(row=r, column=13).value
                p_qty = ws_tn.cell(row=r, column=14).value
                if p_code:
                    po_items.append({"Số Hiệu PO": str(p_code), "Số Lượng Đặt (PCS)": p_qty})
            st.dataframe(pd.DataFrame(po_items), use_container_width=True)
        with po_c2:
            st.markdown("<h5 style='color: #cbd5e1;'>📊 Tổng Hợp Tình Trạng PO</h5>", unsafe_allow_html=True)
            tot_po = safe_num(ws_tn.cell(row=13, column=14).value)
            tot_dlv = safe_num(ws_tn.cell(row=14, column=14).value)
            tot_rem = safe_num(ws_tn.cell(row=15, column=14).value)
            
            st.markdown(f"""
            <div class="metric-card" style="margin-bottom: 10px;">
                <div class="metric-title">Tổng Đơn Đặt Hàng (PO)</div>
                <div class="metric-value">{tot_po:,.0f} PCS</div>
            </div>
            <div class="metric-card">
                <div class="metric-title">Số Lượng Còn Lại Chưa Giao</div>
                <div class="metric-value" style="color: #60a5fa;">{tot_rem:,.0f} PCS</div>
                <div class="metric-sub">Tiến độ đã xuất xưởng: {tot_dlv:,.0f} PCS</div>
            </div>
            """, unsafe_allow_html=True)

# ==================== TAB 3 ====================
with tab_export:
    st.markdown("<h4 style='color: #f1f5f9;'>📥 Trích Xuất & Tải Báo Cáo</h4>", unsafe_allow_html=True)
    with open(target_file, "rb") as f:
        st.download_button(
            label="Tải Xuống Tệp Excel Hoàn Chỉnh (.xlsx)",
            data=f,
            file_name=f"DOSUNG_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
